"""Baut bike-buyers-analysis.xlsx aus der Originaldatei (die unverändert bleibt).

Alles wird über Formeln berechnet (COUNTIFS, AVERAGEIFS, SUMPRODUCT, IF ...), die Filter sind
Dropdowns (Datenüberprüfung). Keine Pivot-Tabellen, keine Datenschnitte, keine Makros.

Aufruf:
    python scripts/build_xlsx.py "Excel Projekt.xlsx" bike-buyers-analysis.xlsx
"""
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SRC = sys.argv[1]
OUT = sys.argv[2]

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
KPI_FILL = PatternFill("solid", fgColor="EAF1FB")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(**kw):
    return Font(name=FONT, **kw)


def header_row(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = font(bold=True)
        c.fill = HEAD_FILL
        c.border = BOX


def title(ws, cell, text, size=14):
    ws[cell] = text
    ws[cell].font = font(bold=True, size=size)


src = load_workbook(SRC)  # nur gelesen, nie gespeichert
raw_src = src["bike_buyers"]
work_src = src["working sheet"]

wb = Workbook()
wb.remove(wb.active)

# ---------------------------------------------------------------- Rohdaten
raw = wb.create_sheet("Rohdaten")
for row in raw_src.iter_rows(values_only=True):
    if row[0] is not None:
        raw.append(list(row))
N_RAW = raw.max_row  # inkl. Überschrift
for c in raw[1]:
    c.font = font(bold=True)
    c.fill = HEAD_FILL
raw.freeze_panes = "A2"

# ---------------------------------------------------------------- Bereinigt
clean = wb.create_sheet("Bereinigt")
headers = ["ID", "Marital Status", "Gender", "Income", "Children", "Education", "Occupation",
           "Home Owner", "Cars", "Commute Distance", "Region", "Age", "Altersgruppe",
           "Purchased Bike", "Kauf (1/0)"]
clean.append(headers)
rows = [list(r) for r in work_src.iter_rows(min_row=2, values_only=True) if r[0] is not None]
for i, vals in enumerate(rows, start=2):
    # Spalten im Original: A..L Daten, M Age Brackets (fehlerhafte Formel), N Purchased Bike
    data = vals[:12]
    data[1] = data[1].strip()  # "Single " -> "Single"
    clean.append(data + [
        f'=IF(L{i}<31,"Unter 31",IF(L{i}<=54,"31–54","55+"))',
        vals[13],
        f'=IF(N{i}="Yes",1,0)',
    ])
LAST = clean.max_row  # 1001
assert LAST == 1001, LAST
for c in clean[1]:
    c.font = font(bold=True)
    c.fill = HEAD_FILL
clean.freeze_panes = "A2"
for col, w in zip("ABCDEFGHIJKLMNO", [8, 14, 9, 10, 9, 18, 15, 11, 6, 16, 14, 6, 13, 14, 11]):
    clean.column_dimensions[col].width = w
for r in range(2, LAST + 1):
    clean[f"D{r}"].number_format = "#,##0"

R = f"Bereinigt!$K$2:$K${LAST}"     # Region
AG = f"Bereinigt!$M$2:$M${LAST}"    # Altersgruppe
P = f"Bereinigt!$N$2:$N${LAST}"     # Purchased Bike
INC = f"Bereinigt!$D$2:$D${LAST}"   # Income
COM = f"Bereinigt!$J$2:$J${LAST}"   # Commute Distance
FILT = f'{R},$C$4,{AG},$C$5'        # aktive Dashboard-Filter (über Wildcard "*" für "Alle")

# ---------------------------------------------------------------- Datenprüfung
chk = wb.create_sheet("Datenprüfung")
title(chk, "A1", "Datenprüfung und Bereinigung")
header_row(chk, 3, ["Prüfung", "Rohdaten", "Bereinigt"])
RL = N_RAW
checks = [
    ("Datensätze", f"=COUNTA(Rohdaten!A2:A{RL})", f"=COUNTA(Bereinigt!A2:A{LAST})"),
    ("Eindeutige IDs", f"=SUMPRODUCT(1/COUNTIF(Rohdaten!A2:A{RL},Rohdaten!A2:A{RL}))",
     f"=SUMPRODUCT(1/COUNTIF(Bereinigt!A2:A{LAST},Bereinigt!A2:A{LAST}))"),
    ("Doppelte Datensätze", "=B4-B5", "=C4-C5"),
    ("Leere Zellen", f"=COUNTBLANK(Rohdaten!A2:M{RL})", f"=COUNTBLANK(Bereinigt!A2:O{LAST})"),
    ("Alter min", f"=MIN(Rohdaten!L2:L{RL})", f"=MIN(Bereinigt!L2:L{LAST})"),
    ("Alter max", f"=MAX(Rohdaten!L2:L{RL})", f"=MAX(Bereinigt!L2:L{LAST})"),
    ("Einkommen min", f"=MIN(Rohdaten!D2:D{RL})", f"=MIN(Bereinigt!D2:D{LAST})"),
    ("Einkommen max", f"=MAX(Rohdaten!D2:D{RL})", f"=MAX(Bereinigt!D2:D{LAST})"),
    ("Familienstand mit Leerzeichen am Rand",
     f"=SUMPRODUCT(--(LEN(Rohdaten!B2:B{RL})<>LEN(TRIM(Rohdaten!B2:B{RL}))))",
     f"=SUMPRODUCT(--(LEN(Bereinigt!B2:B{LAST})<>LEN(TRIM(Bereinigt!B2:B{LAST}))))"),
    ("Personen ohne gültige Altersgruppe", "–",
     f'=COUNTA(Bereinigt!M2:M{LAST})-COUNTIF(Bereinigt!M2:M{LAST},"Unter 31")'
     f'-COUNTIF(Bereinigt!M2:M{LAST},"31–54")-COUNTIF(Bereinigt!M2:M{LAST},"55+")'),
]
for i, (name, a, b) in enumerate(checks, start=4):
    chk.cell(row=i, column=1, value=name)
    chk.cell(row=i, column=2, value=a)
    chk.cell(row=i, column=3, value=b)
    for c in (2, 3):
        chk.cell(row=i, column=c).number_format = "#,##0"
        chk.cell(row=i, column=c).alignment = Alignment(horizontal="right")
r = 4 + len(checks) + 1
chk.cell(row=r, column=1, value="Verschiedene Werte je Kategorie (bereinigt)").font = font(bold=True)
r += 1
header_row(chk, r, ["Spalte", "", "Anzahl Werte"])
r += 1
for name, col in [("Marital Status", "B"), ("Gender", "C"), ("Education", "F"), ("Occupation", "G"),
                  ("Home Owner", "H"), ("Commute Distance", "J"), ("Region", "K"),
                  ("Altersgruppe", "M"), ("Purchased Bike", "N")]:
    chk.cell(row=r, column=1, value=name)
    chk.cell(row=r, column=3, value=f"=SUMPRODUCT(1/COUNTIF(Bereinigt!{col}2:{col}{LAST},Bereinigt!{col}2:{col}{LAST}))")
    chk.cell(row=r, column=3).number_format = "0"
    r += 1
r += 1
chk.cell(row=r, column=1, value="Änderungsprotokoll").font = font(bold=True)
r += 1
for line in [
    "1. 26 doppelte Datensätze entfernt (komplett identische Zeilen mit gleicher ID), es bleiben 1.000 Personen. (aus dem Tutorial)",
    "2. Familienstand M/S zu Married/Single und Geschlecht M/F zu Male/Female ausgeschrieben. (aus dem Tutorial)",
    "3. Leerzeichen am Ende von „Single “ entfernt. (Überarbeitung)",
    "4. Spaltenüberschrift „Mariatl Statu “ zu „Marital Status“ korrigiert. (Überarbeitung)",
    "5. Altersgruppen neu: Unter 31 / 31–54 / 55+. In der alten Formel landeten alle unter 31 in „Invalid“. (Überarbeitung)",
    "6. Hilfsspalte „Kauf (1/0)“ ergänzt. (Überarbeitung)",
    "Das Blatt „Rohdaten“ ist unverändert und entspricht der Originaldatei.",
]:
    chk.cell(row=r, column=1, value=line)
    r += 1
chk.column_dimensions["A"].width = 44
chk.column_dimensions["B"].width = 14
chk.column_dimensions["C"].width = 14

# ---------------------------------------------------------------- Auswertung
ana = wb.create_sheet("Auswertung")
title(ana, "A1", "Auswertung (alle Werte per Formel, gesteuert über die Filter im Dashboard)", 13)

ana["A3"] = "Aktiver Filter"; ana["A3"].font = font(bold=True)
ana["B3"] = "Auswahl"; ana["B3"].font = font(bold=True)
ana["C3"] = "Kriterium"; ana["C3"].font = font(bold=True)
ana["A4"] = "Region"
ana["B4"] = "=Dashboard!$C$4"
ana["C4"] = '=IF(B4="Alle","*",B4)'
ana["A5"] = "Altersgruppe"
ana["B5"] = "=Dashboard!$F$4"
ana["C5"] = '=IF(B5="Alle","*",B5)'
ana["E3"] = '"*" steht für alle Werte (Platzhalter in COUNTIFS).'
ana["E3"].font = font(italic=True, color="808080")

# Auswahllisten für die Dropdowns
ana["K3"] = "Liste Region"; ana["K3"].font = font(bold=True)
for i, v in enumerate(["Alle", "Europe", "North America", "Pacific"], start=4):
    ana[f"K{i}"] = v
ana["L3"] = "Liste Altersgruppe"; ana["L3"].font = font(bold=True)
for i, v in enumerate(["Alle", "Unter 31", "31–54", "55+"], start=4):
    ana[f"L{i}"] = v

# Kennzahlen
title(ana, "A7", "Kennzahlen", 11)
ana["A8"] = "Personen"; ana["B8"] = f"=COUNTIFS({FILT})"
ana["A9"] = "Käufer"; ana["B9"] = f'=COUNTIFS({FILT},{P},"Yes")'
ana["A10"] = "Kaufquote"; ana["B10"] = "=IF(B8=0,0,B9/B8)"
ana["B8"].number_format = ana["B9"].number_format = "#,##0"
ana["B10"].number_format = "0.0%"

# Altersgruppen (nur Regionsfilter)
title(ana, "A12", "Kaufquote nach Altersgruppe (nur Regionsfilter)", 11)
header_row(ana, 13, ["Altersgruppe", "Personen", "Käufer", "Kaufquote", "Beschriftung"])
for i, g in enumerate(["Unter 31", "31–54", "55+"], start=14):
    ana[f"A{i}"] = g
    ana[f"B{i}"] = f"=COUNTIFS({R},$C$4,{AG},A{i})"
    ana[f"C{i}"] = f'=COUNTIFS({R},$C$4,{AG},A{i},{P},"Yes")'
    ana[f"D{i}"] = f"=IF(B{i}=0,0,C{i}/B{i})"
    ana[f"E{i}"] = f'=A{i}&" (n="&B{i}&")"'
ana["A17"] = "Gesamt"; ana["B17"] = "=SUM(B14:B16)"; ana["C17"] = "=SUM(C14:C16)"; ana["D17"] = "=IF(B17=0,0,C17/B17)"
AGE_ROWS = (14, 16)

# Pendelstrecke (beide Filter)
title(ana, "A19", "Kaufquote nach Pendelstrecke", 11)
header_row(ana, 20, ["Pendelstrecke", "Personen", "Käufer", "Kaufquote", "Beschriftung"])
for i, g in enumerate(["0-1 Miles", "1-2 Miles", "2-5 Miles", "5-10 Miles", "10+ Miles"], start=21):
    ana[f"A{i}"] = g
    ana[f"B{i}"] = f"=COUNTIFS({FILT},{COM},A{i})"
    ana[f"C{i}"] = f'=COUNTIFS({FILT},{COM},A{i},{P},"Yes")'
    ana[f"D{i}"] = f"=IF(B{i}=0,0,C{i}/B{i})"
    ana[f"E{i}"] = f'=A{i}&" (n="&B{i}&")"'
ana["A26"] = "Gesamt"; ana["B26"] = "=SUM(B21:B25)"; ana["C26"] = "=SUM(C21:C25)"; ana["D26"] = "=IF(B26=0,0,C26/B26)"
COM_ROWS = (21, 25)

# Einkommen Käufer vs. Nichtkäufer (beide Filter)
title(ana, "A28", "Einkommen: Käufer vs. Nichtkäufer", 11)
header_row(ana, 29, ["Gruppe", "Personen", "Ø Einkommen"])
for i, (g, v) in enumerate([("Käufer", "Yes"), ("Nichtkäufer", "No")], start=30):
    ana[f"A{i}"] = g
    ana[f"B{i}"] = f'=COUNTIFS({FILT},{P},"{v}")'
    ana[f"C{i}"] = f'=IF(B{i}=0,0,AVERAGEIFS({INC},{FILT},{P},"{v}"))'
ana["A32"] = "Alle"; ana["B32"] = "=B30+B31"
ana["C32"] = f"=IF(B32=0,0,AVERAGEIFS({INC},{FILT}))"

title(ana, "A34", "Verteilung nach Einkommensklasse (Anteil innerhalb der Gruppe)", 11)
header_row(ana, 35, ["Einkommensklasse", "von", "bis", "Käufer", "Nichtkäufer", "Anteil Käufer", "Anteil Nichtkäufer"])
bands = [("unter 40.000", 0, 39999), ("40.000–69.999", 40000, 69999),
         ("70.000–99.999", 70000, 99999), ("ab 100.000", 100000, 10000000)]
for i, (g, lo, hi) in enumerate(bands, start=36):
    ana[f"A{i}"] = g
    ana[f"B{i}"] = lo
    ana[f"C{i}"] = hi
    ana[f"D{i}"] = f'=COUNTIFS({FILT},{P},"Yes",{INC},">="&B{i},{INC},"<="&C{i})'
    ana[f"E{i}"] = f'=COUNTIFS({FILT},{P},"No",{INC},">="&B{i},{INC},"<="&C{i})'
    ana[f"F{i}"] = f"=IF($B$30=0,0,D{i}/$B$30)"
    ana[f"G{i}"] = f"=IF($B$31=0,0,E{i}/$B$31)"
ana["A40"] = "Summe"; ana["D40"] = "=SUM(D36:D39)"; ana["E40"] = "=SUM(E36:E39)"
ana["F40"] = "=SUM(F36:F39)"; ana["G40"] = "=SUM(G36:G39)"
INC_ROWS = (36, 39)

for row in ana.iter_rows(min_row=1, max_row=40):
    for c in row:
        if c.font is None or c.font.name != FONT:
            c.font = Font(name=FONT, bold=c.font.b if c.font else False,
                          italic=c.font.i if c.font else False,
                          size=c.font.sz if c.font and c.font.sz else 10,
                          color=c.font.color if c.font else None)
for r in list(range(14, 18)) + list(range(21, 27)):
    ana[f"B{r}"].number_format = ana[f"C{r}"].number_format = "#,##0"
    ana[f"D{r}"].number_format = "0.0%"
for r in range(30, 33):
    ana[f"B{r}"].number_format = "#,##0"
    ana[f"C{r}"].number_format = "#,##0"
for r in range(36, 41):
    for col in "BCDE":
        ana[f"{col}{r}"].number_format = "#,##0"
    ana[f"F{r}"].number_format = ana[f"G{r}"].number_format = "0.0%"
for col, w in zip("ABCDEFGHIJKL", [24, 14, 14, 12, 22, 14, 18, 2, 2, 2, 15, 18]):
    ana.column_dimensions[col].width = w

# ---------------------------------------------------------------- Dashboard
dash = wb.create_sheet("Dashboard")
wb.move_sheet("Dashboard", offset=-(len(wb.sheetnames) - 1))  # als erstes Blatt
wb.active = 0
dash.sheet_view.showGridLines = False
for col in "ABCDEFGHIJKLMNOPQ":
    dash.column_dimensions[col].width = 11
dash.column_dimensions["A"].width = 2
dash.column_dimensions["E"].width = 14
dash.page_setup.orientation = "landscape"
dash.page_setup.fitToWidth = 1
dash.page_setup.fitToHeight = 1
dash.sheet_properties.pageSetUpPr.fitToPage = True

dash["B1"] = "Fahrradkäufer – Dashboard"
dash["B1"].font = font(bold=True, size=18)
dash["B2"] = ("Datenbasis: 1.000 Personen (Bike Buyers, bereinigt). Filter über die gelben Dropdowns. "
              "Alle Werte werden per Formel berechnet.")
dash["B2"].font = font(italic=True, color="808080")

dash["B4"] = "Region:"; dash["B4"].font = font(bold=True); dash["B4"].alignment = Alignment(horizontal="right")
dash["C4"] = "Alle"
dash["E4"] = "Altersgruppe:"; dash["E4"].font = font(bold=True); dash["E4"].alignment = Alignment(horizontal="right")
dash["F4"] = "Alle"
for ref in ("C4", "F4"):
    dash[ref].fill = INPUT_FILL
    dash[ref].border = BOX
    dash[ref].font = font(bold=True)
dash.merge_cells("C4:D4")
dash.merge_cells("F4:G4")

dv_region = DataValidation(type="list", formula1="Auswertung!$K$4:$K$7", allow_blank=False,
                           showErrorMessage=True, errorTitle="Ungültige Auswahl",
                           error="Bitte einen Wert aus der Liste wählen.")
dv_age = DataValidation(type="list", formula1="Auswertung!$L$4:$L$7", allow_blank=False,
                        showErrorMessage=True, errorTitle="Ungültige Auswahl",
                        error="Bitte einen Wert aus der Liste wählen.")
dash.add_data_validation(dv_region)
dash.add_data_validation(dv_age)
dv_region.add("C4")
dv_age.add("F4")

kpis = [("B", "C", "Personen", "=Auswertung!B8", "#,##0"),
        ("E", "F", "Käufer", "=Auswertung!B9", "#,##0"),
        ("H", "I", "Kaufquote", "=Auswertung!B10", "0.0%")]
for c1, c2, label, formula, fmt in kpis:
    dash.merge_cells(f"{c1}6:{c2}6")
    dash.merge_cells(f"{c1}7:{c2}7")
    dash[f"{c1}6"] = label
    dash[f"{c1}6"].font = font(color="595959")
    dash[f"{c1}7"] = formula
    dash[f"{c1}7"].font = font(bold=True, size=20)
    dash[f"{c1}7"].number_format = fmt
    for rr in (6, 7):
        for cc in (c1, c2):
            dash[f"{cc}{rr}"].fill = KPI_FILL
            dash[f"{cc}{rr}"].border = BOX
            dash[f"{cc}{rr}"].alignment = Alignment(horizontal="center", vertical="center")
dash.row_dimensions[7].height = 32

dash["K6"] = "Hinweise"
dash["K6"].font = font(bold=True)
for i, t in enumerate([
    "n = Anzahl Personen in der Gruppe (bei kleinem n vorsichtig lesen).",
    "Das Altersdiagramm reagiert nur auf den Regionsfilter.",
    "Gezeigt werden Zusammenhänge, keine Ursachen.",
], start=7):
    dash[f"K{i}"] = t
    dash[f"K{i}"].font = font(size=9, color="595959")


def pct_labels(chart):
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.showLegendKey = False
    chart.dataLabels.showPercent = False
    chart.dataLabels.numFmt = "0%"


def style_axes(chart, y_fmt):
    chart.y_axis.number_format = y_fmt
    chart.y_axis.majorGridlines = None
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.scaling.min = 0


# Diagramm 1: Kaufquote nach Altersgruppe
c1 = BarChart(); c1.type = "col"
c1.title = "Kaufquote nach Altersgruppe"
c1.add_data(Reference(ana, min_col=4, min_row=AGE_ROWS[0], max_row=AGE_ROWS[1]), titles_from_data=False)
c1.set_categories(Reference(ana, min_col=5, min_row=AGE_ROWS[0], max_row=AGE_ROWS[1]))
c1.series[0].tx = SeriesLabel(v="Kaufquote")
c1.legend = None
style_axes(c1, "0%")
c1.y_axis.scaling.max = 1
pct_labels(c1)
c1.width, c1.height = 15, 7.5
dash.add_chart(c1, "B10")

# Diagramm 2: Kaufquote nach Pendelstrecke
c2 = BarChart(); c2.type = "col"
c2.title = "Kaufquote nach Pendelstrecke"
c2.add_data(Reference(ana, min_col=4, min_row=COM_ROWS[0], max_row=COM_ROWS[1]), titles_from_data=False)
c2.set_categories(Reference(ana, min_col=5, min_row=COM_ROWS[0], max_row=COM_ROWS[1]))
c2.series[0].tx = SeriesLabel(v="Kaufquote")
c2.legend = None
style_axes(c2, "0%")
c2.y_axis.scaling.max = 1
pct_labels(c2)
c2.width, c2.height = 15, 7.5
dash.add_chart(c2, "I10")

# Diagramm 3: Einkommensverteilung Käufer vs. Nichtkäufer
c3 = BarChart(); c3.type = "col"; c3.grouping = "clustered"
c3.title = "Einkommen: Käufer vs. Nichtkäufer (Anteil je Gruppe)"
c3.add_data(Reference(ana, min_col=6, max_col=7, min_row=35, max_row=INC_ROWS[1]), titles_from_data=True)
c3.set_categories(Reference(ana, min_col=1, min_row=INC_ROWS[0], max_row=INC_ROWS[1]))
c3.legend.position = "b"
style_axes(c3, "0%")
pct_labels(c3)
c3.width, c3.height = 15, 7.5
dash.add_chart(c3, "B26")

# Ø Einkommen als kleine Tabelle neben Diagramm 3
dash["I26"] = "Ø Einkommen"
dash["I26"].font = font(bold=True)
for i, (label, ref) in enumerate([("Käufer", "C30"), ("Nichtkäufer", "C31"), ("Alle", "C32")], start=27):
    dash[f"I{i}"] = label
    dash[f"I{i}"].font = font()
    dash[f"K{i}"] = f"=Auswertung!{ref}"
    dash[f"K{i}"].number_format = "#,##0"
    dash[f"K{i}"].font = font(bold=True)

# Schrift überall auf Arial
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                f = c.font
                c.font = Font(name=FONT, bold=f.b, italic=f.i, size=f.sz, color=f.color)

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("gespeichert:", OUT, wb.sheetnames)
